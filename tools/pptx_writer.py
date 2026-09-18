#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
.pptx 生成器（基于 python-pptx）
===============================
把结构化的幻灯内容写成真正的 PowerPoint 文件，供开题报告 / 答辩汇报使用。

为什么用 python-pptx：本包已经要求 matplotlib + numpy 才能出图（见 requirements.txt），
"零依赖"从来只对纯统计脚本成立；PPT 属于"产出物"这一类，和图表同级。
需要零安装的手机版是单文件形态，那种形态下统计分析与 PPT 生成本来就做不了，
所以这里不迁就它。

依赖缺失时**优雅降级**：`available()` 返回 False，调用方应改为输出 markdown 大纲
并提示一行安装命令，而不是崩在 ImportError 上。

版式用 python-pptx 默认模板的 layout 名，按名字取而不是按序号，避免模板版本变化时错位：
  Title Slide / Title and Content / Section Header / Title Only / Blank

页面尺寸：默认 4:3（默认模板的占位符坐标按 4:3 排布）。
ratio="16:9" 时会把母版与全部版式里显式定位的形状横坐标按 4/3 缩放，
所以内容仍然是居中的，不会挤在左边。
"""

_LAYOUT_NAMES = {
    "title": "Title Slide",
    "bullets": "Title and Content",
    "section": "Section Header",
    "title_only": "Title Only",
    "blank": "Blank",
}

SLIDE_W_4X3 = 9144000     # EMU，10in
SLIDE_H = 6858000         # 7.5in

# --- 输出编码守卫：管道/重定向时强制 UTF-8（项目门禁统一要求）---
import sys as _sys
if hasattr(_sys.stdout, "reconfigure") and not _sys.stdout.isatty():
    _sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    _sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def _pptx():
    """延迟导入：没装 python-pptx 时也能 import 本模块、能走降级分支。"""
    try:
        import pptx  # noqa: F401
        from pptx import Presentation  # noqa: F401
        from pptx.util import Inches, Pt, Emu  # noqa: F401
    except ImportError:
        return None
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    return {"Presentation": Presentation, "Inches": Inches, "Pt": Pt, "Emu": Emu}


def available():
    """python-pptx 是否就绪（调用方据此决定生成 .pptx 还是退回 markdown）。"""
    return _pptx() is not None


INSTALL_HINT = "pip install python-pptx"


class Presentation:
    """攒一页页内容，最后 save() 成 .pptx。"""

    def __init__(self, title="", author="", ratio="4:3"):
        api = _pptx()
        if api is None:
            raise RuntimeError(
                f"生成 .pptx 需要 python-pptx，请先执行：{INSTALL_HINT}\n"
                "（未安装时请改用 markdown 大纲交付，不要让流程断在这里。）")
        self._api = api
        self.prs = api["Presentation"]()
        self.prs.slide_width = api["Emu"](SLIDE_W_4X3)
        self.prs.slide_height = api["Emu"](SLIDE_H)
        self.ratio = ratio
        # 必须在加任何幻灯片之前改尺寸：形状是按当前 slide_width 定位的，
        # 事后再放大页面会把先放的内容留在左边。
        if ratio == "16:9":
            self._to_169()
        core = self.prs.core_properties
        if title:
            core.title = title
        if author:
            core.author = core.last_modified_by = author
        self._count = 0

    # ---------------- 内部 ----------------
    def _layout(self, key):
        want = _LAYOUT_NAMES[key]
        for lay in self.prs.slide_layouts:
            if lay.name == want:
                return lay
        return self.prs.slide_layouts[5]      # Title Only 兜底

    def _slide(self, key):
        s = self.prs.slides.add_slide(self._layout(key))
        self._count += 1
        return s

    @staticmethod
    def _set(tf, text, size=None, bold=False):
        """写文本；size 单位是磅。"""
        api = _pptx()
        tf.text = text
        p = tf.paragraphs[0]
        if size:
            p.font.size = api["Pt"](size)
        if bold:
            p.font.bold = True
        return p

    def _notes(self, slide, text):
        if text:
            self._set(slide.notes_slide.notes_text_frame, text, size=12)

    # ---------------- 公开：各种页 ----------------
    def add_title(self, title, subtitle="", notes=""):
        s = self._slide("title")
        self._set(s.shapes[0].text_frame, title, size=40, bold=True)
        if subtitle and len(s.shapes) > 1:
            self._set(s.shapes[1].text_frame, subtitle, size=18)
        self._notes(s, notes)
        return self._count - 1

    def add_bullets(self, title, items, notes=""):
        """items 元素可以是 str、(level, text) 或只有 1 个元素的元组。"""
        s = self._slide("bullets")
        self._set(s.shapes[0].text_frame, title, size=30, bold=True)
        body = self._body(s)
        if body is None:
            self._notes(s, notes)
            return self._count - 1
        first = True
        for it in items:
            if isinstance(it, (tuple, list)):
                lvl, text = (it[0], it[1]) if len(it) >= 2 else (0, it[0])
            else:
                lvl, text = 0, it
            try:
                lvl = min(max(int(lvl), 0), 4)
            except (TypeError, ValueError):
                lvl, text = 0, f"{lvl}{text}"
            p = body.paragraphs[0] if first else body.add_paragraph()
            first = False
            p.text = str(text)
            p.level = lvl
            p.font.size = self._api["Pt"](22 if lvl == 0 else 18)
        self._notes(s, notes)
        return self._count - 1

    def add_table(self, title, header, rows, notes="", font_sz=14):
        s = self._slide("title_only")
        self._set(s.shapes[0].text_frame, title, size=30, bold=True)
        ncols = max([len(header or [])] + [len(r) for r in rows] + [1])
        nrows = len(rows) + (1 if header else 0)
        if nrows and ncols > 1:
            x = self._api["Inches"](0.7)
            cx = int(self.prs.slide_width - 2 * x)
            y = int(SLIDE_H * 0.26)
            cy = min(int(SLIDE_H * 0.62), self._api["Inches"](0.42) * nrows)
            gf = s.shapes.add_table(nrows, ncols, x, y, cx, cy)
            t = gf.table
            r0 = 0
            if header:
                for c, v in enumerate(header):
                    self._cell(t.cell(0, c), v, font_sz, bold=True)
                r0 = 1
            for i, row in enumerate(rows):
                for c in range(ncols):
                    v = row[c] if c < len(row) else ""
                    self._cell(t.cell(r0 + i, c), v, font_sz)
        self._notes(s, notes)
        return self._count - 1

    def add_picture(self, title, png_path, caption="", notes=""):
        from pathlib import Path
        p = Path(png_path)
        if not p.exists():
            raise FileNotFoundError(f"图片不存在：{p}")
        s = self._slide("title_only")
        self._set(s.shapes[0].text_frame, title, size=30, bold=True)
        w, h = _png_size(p)
        box_w = int(self.prs.slide_width * 0.72)
        box_h = int(SLIDE_H * 0.58)
        scale = min(box_w / max(w, 1), box_h / max(h, 1))
        dw, dh = int(w * scale), int(h * scale)
        left = int((self.prs.slide_width - dw) / 2)
        top = int(SLIDE_H * 0.22)
        s.shapes.add_picture(str(p), self._api["Emu"](left), self._api["Emu"](top),
                             width=self._api["Emu"](dw), height=self._api["Emu"](dh))
        if caption:
            tb = s.shapes.add_textbox(self._api["Emu"](left),
                                      self._api["Emu"](top + dh + 60000),
                                      self._api["Emu"](dw), self._api["Emu"](320000))
            self._set(tb.text_frame, caption, size=12)
        self._notes(s, notes)
        return self._count - 1

    def add_section(self, title, notes=""):
        s = self._slide("section")
        self._set(s.shapes[0].text_frame, title, size=32, bold=True)
        self._notes(s, notes)
        return self._count - 1

    def add_free(self, title, lines, notes=""):
        """仅标题 + 若干自由文本行（不依赖正文占位符，版式最不容易出问题）。"""
        s = self._slide("title_only")
        self._set(s.shapes[0].text_frame, title, size=30, bold=True)
        tb = s.shapes.add_textbox(self._api["Inches"](0.8), self._api["Inches"](2.0),
                                  int(self.prs.slide_width - 2 * self._api["Inches"](0.8)),
                                  int(SLIDE_H * 0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = str(line)
            p.font.size = self._api["Pt"](20)
        self._notes(s, notes)
        return self._count - 1

    # ---------------- 落盘 ----------------
    def _body(self, slide):
        """取正文占位符；模板缺该占位符时返回 None 而不是抛异常。"""
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                return ph.text_frame
        return None

    def _cell(self, cell, text, size, bold=False):
        self._set(cell.text_frame, str(text), size=size, bold=bold)

    def save(self, path):
        from pathlib import Path
        out = Path(path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return out

    def _to_169(self):
        """切 16:9：页面加宽 4/3，并把母版/版式里**自己写了坐标**的形状横坐标同比放大。

        只动显式 `<a:xfrm>`，绝不碰纯继承的形状：版式占位符默认继承母版，
        读 `sh.left` 会拿到继承来的值，一旦写回去就变成"母版已放大 × 再放大一次"=
        实际缩放 (4/3)²，占位符会顶出右边界（这个坑踩过，实测超出 3.56 英寸）。
        """
        from pptx.oxml.ns import qn
        new_w = int(SLIDE_W_4X3 * 4 / 3)
        k = new_w / SLIDE_W_4X3
        self.prs.slide_width = self._api["Emu"](new_w)
        touched = 0
        for master in self.prs.slide_masters:
            holders = [master] + list(master.slide_layouts)
            for holder in holders:
                for sh in holder.shapes:
                    spPr = sh._element.find(qn("p:spPr"))
                    if spPr is None:
                        continue
                    xfrm = spPr.find(qn("a:xfrm"))
                    if xfrm is None:
                        continue          # 没有显式几何 = 继承，跳过，不要写死
                    off = xfrm.find(qn("a:off"))
                    ext = xfrm.find(qn("a:ext"))
                    if off is not None and off.get("x") is not None:
                        off.set("x", str(int(int(off.get("x")) * k)))
                        touched += 1
                    if ext is not None and ext.get("cx") is not None:
                        ext.set("cx", str(int(int(ext.get("cx")) * k)))
        self._scaled = touched

    @property
    def scaled_shapes(self):
        """16:9 下被显式重定位的形状个数（0 说明模板全靠继承，页面加宽也不会错位）。"""
        return getattr(self, "_scaled", 0)


def _png_size(path):
    """从 PNG IHDR 读宽高；非 PNG 抛错，避免写出坏图。"""
    with open(path, "rb") as f:
        head = f.read(26)
    if head[:8] != b"\x89PNG\r\n\x1a\n" or head[12:16] != b"IHDR":
        raise ValueError(f"只支持 PNG：{path}")
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")
